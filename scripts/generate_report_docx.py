# -*- coding: utf-8 -*-
"""生成大作业 Word 报告（运行：python scripts/generate_report_docx.py）"""
from pathlib import Path

from docx import Document
from docx.enum.text import WD_LINE_SPACING
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs" / "大作业-Offer捕手.docx"

FONT_HEI = "黑体"
FONT_SONG = "宋体"
SIZE_H1 = Pt(16)   # 三号
SIZE_H2 = Pt(14)   # 四号
SIZE_H3 = Pt(10.5)  # 五号
SIZE_BODY = Pt(10.5)
LINE_SPACING = Pt(20)


def set_run_font(run, name: str, size: Pt, bold: bool = False):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = size
    run.font.bold = bold
    run.font.color.rgb = RGBColor(0, 0, 0)


def set_paragraph_format(paragraph, first_line_indent: bool = True):
    pf = paragraph.paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    pf.line_spacing = LINE_SPACING
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)
    if first_line_indent:
        pf.first_line_indent = Pt(21)  # 约 2 字符


def add_heading1(doc: Document, text: str):
    p = doc.add_paragraph()
    run = p.add_run(text)
    set_run_font(run, FONT_HEI, SIZE_H1, bold=True)
    set_paragraph_format(p, first_line_indent=False)
    return p


def add_heading2(doc: Document, text: str):
    p = doc.add_paragraph()
    run = p.add_run(text)
    set_run_font(run, FONT_HEI, SIZE_H2, bold=True)
    set_paragraph_format(p, first_line_indent=False)
    return p


def add_heading3(doc: Document, text: str):
    p = doc.add_paragraph()
    run = p.add_run(text)
    set_run_font(run, FONT_SONG, SIZE_H3, bold=True)
    set_paragraph_format(p, first_line_indent=False)
    return p


def add_body(doc: Document, text: str):
    p = doc.add_paragraph()
    run = p.add_run(text)
    set_run_font(run, FONT_SONG, SIZE_BODY)
    set_paragraph_format(p, first_line_indent=True)
    return p


def add_screenshot_placeholder(doc: Document, caption: str):
    p = doc.add_paragraph()
    run = p.add_run(f"【此处插入截图：{caption}】")
    set_run_font(run, FONT_SONG, SIZE_BODY)
    run.font.italic = True
    run.font.color.rgb = RGBColor(128, 128, 128)
    set_paragraph_format(p, first_line_indent=False)
    p.paragraph_format.space_after = Pt(6)
    # 预留空白行便于粘贴图片
    blank = doc.add_paragraph()
    blank.add_run("\n")
    set_paragraph_format(blank, first_line_indent=False)


def build():
    doc = Document()
    # 正文默认样式
    style = doc.styles["Normal"]
    style.font.name = FONT_SONG
    style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_SONG)
    style.font.size = SIZE_BODY

    # 封面信息（可按学校模板调整）
    t = doc.add_paragraph()
    r = t.add_run("课程大作业报告")
    set_run_font(r, FONT_HEI, SIZE_H1, bold=True)
    set_paragraph_format(t, first_line_indent=False)
    t.alignment = 1  # center

    for line in [
        "",
        "作品名称：Offer 捕手——学生求职匹配智能体",
        "姓    名：（请填写）",
        "学    号：（请填写）",
        "专    业：（请填写）",
        "指导教师：（请填写）",
        "完成日期：（请填写）",
    ]:
        p = doc.add_paragraph()
        run = p.add_run(line)
        set_run_font(run, FONT_SONG, SIZE_BODY)
        set_paragraph_format(p, first_line_indent=False)
        p.alignment = 1

    doc.add_page_break()

  # ========== 一、作品背景 ==========
    add_heading1(doc, "一、作品背景")

    add_body(
        doc,
        "随着高校毕业生规模持续扩大，在校大学生与应届生面临实习、校招双线并进的求职压力。"
        "传统求职方式依赖招聘网站人工筛选岗位、反复修改简历，存在信息分散、匹配效率低、"
        "难以判断自身与目标岗位差距等问题。尤其在互联网、金融、制造等行业，岗位描述专业性强，"
        "学生往往不清楚应突出哪些经历与技能。",
    )
    add_body(
        doc,
        "近年来，大语言模型（LLM）在文本理解、结构化抽取与推理方面能力显著增强，"
        "为“简历—岗位”语义匹配、个性化建议与自然语言咨询提供了新的技术路径。"
        "在此背景下，本课程大作业设计并实现 Web 应用「Offer 捕手」："
        "面向学生用户提供岗位库管理、简历解析、智能匹配分析、简历优化建议及 AI 对话等功能，"
        "力求在可部署、可演示的前提下，形成一条完整的“注册—画像—简历—岗位—分析”业务闭环。",
    )
    add_screenshot_placeholder(doc, "图1-1 系统首页/对话主界面整体效果")

    add_heading2(doc, "1.1 需求来源与痛点")
    add_body(
        doc,
        "通过对身边同学求职过程的访谈与归纳，主要痛点包括：（1）岗位搜寻耗时，"
        "需在多个平台复制 JD；（2）简历与岗位匹配度缺乏量化反馈；（3）不知如何针对具体岗位优化表述；"
        "（4）零散问题缺少可追问的咨询入口。本作品围绕上述痛点进行功能拆解与实现。",
    )

    add_heading2(doc, "1.2 目标用户")
    add_body(
        doc,
        "系统主要服务对象为在校大学生、应届毕业生，求职类型以实习与校招为主。"
        "用户具备基本 Web 使用能力，可通过邮箱注册并上传 PDF/Word 格式简历。",
    )

    # ========== 二、作品目标 ==========
    add_heading1(doc, "二、作品目标")

    add_heading2(doc, "2.1 总体目标")
    add_body(
        doc,
        "构建一套前后端分离的求职辅助 Web 系统，实现用户认证、求职意向管理、简历上传与解析、"
        "岗位库维护、基于规则与大模型结合的匹配推荐、单岗位深度分析报告、"
        "针对岗位的简历优化建议，以及基于 DeepSeek 的智能对话模块；"
        "支持 Docker Compose 一键部署至本机或云服务器，便于演示与验收。",
    )

    add_heading2(doc, "2.2 具体功能目标")
    add_body(doc, "（1）用户与认证：邮箱注册，邮件链接设置密码，JWT 登录；登录失败与发信频率限流。")
    add_body(doc, "（2）求职画像：维护目标城市、岗位方向、行业及实习/校招类型，供匹配模块使用。")
    add_body(doc, "（3）简历管理：支持 PDF、DOCX、TXT 上传，解析为结构化数据，支持多版本与当前简历切换。")
    add_body(doc, "（4）岗位库：预置种子岗位；支持手动录入、粘贴/截图解析心仪岗位、AI 检索导入。")
    add_body(doc, "（5）匹配与优化：岗位库智能推荐 Top-K；单岗匹配报告；简历优化建议（不编造经历）。")
    add_body(doc, "（6）智能对话：DeepSeek 风格主界面，支持多会话历史与用户自备 API Key。")

    add_screenshot_placeholder(doc, "图2-1 功能模块关系或导航结构示意（可用 draw.io / 架构简图）")

    add_heading2(doc, "2.3 非功能目标")
    add_body(doc, "安全性：密码 bcrypt 存储，接口 JWT 鉴权，敏感配置写入 .env 不入库。")
    add_body(doc, "性能：热点列表接口使用 Redis 缓存；登录限流防止暴力尝试。")
    add_body(doc, "可维护性：REST API 与 OpenAPI 文档；模块化划分 api / services / models。")

    # ========== 三、技术选型 ==========
    add_heading1(doc, "三、技术选型")

    add_heading2(doc, "3.1 总体架构")
    add_body(
        doc,
        "系统采用 B/S 架构与前后端分离设计。浏览器访问 Nginx 托管的 React 单页应用；"
        "前端通过 /api 反向代理访问 FastAPI 后端；后端连接 PostgreSQL 持久化业务数据、"
        "Redis 承担限流与缓存，并通过 HTTPS 调用 DeepSeek 开放 API 完成智能分析。",
    )
    add_screenshot_placeholder(doc, "图3-1 系统架构图（可参考 docs/architecture.md 中 mermaid 导出）")

    add_heading2(doc, "3.2 前端技术")
    add_heading3(doc, "3.2.1 框架与构建")
    add_body(doc, "选用 React 18 + TypeScript，构建工具为 Vite，样式使用 Tailwind CSS，路由采用 React Router。")
    add_body(doc, "选型理由：组件化开发效率高，TypeScript 利于接口类型约束，Vite 热更新与构建速度快。")

    add_heading3(doc, "3.2.2 交互与展示")
    add_body(doc, "对话内容支持 Markdown 渲染；岗位、匹配等页面采用卡片式布局；认证页采用独立 AuthLayout 组件统一风格。")
    add_screenshot_placeholder(doc, "图3-2 前端页面示例（登录页或岗位库页）")

    add_heading2(doc, "3.3 后端技术")
    add_heading3(doc, "3.3.1 Web 框架与 ORM")
    add_body(doc, "后端基于 Python 3.11 与 FastAPI，异步数据库访问使用 SQLAlchemy 2.0 + asyncpg 驱动 PostgreSQL。")
    add_body(doc, "请求/响应模型使用 Pydantic 校验；自动生成 Swagger 文档便于联调。")

    add_heading3(doc, "3.3.2 认证与安全")
    add_body(doc, "密码哈希采用 passlib+bcrypt；会话采用无状态 JWT（HS256）；邮件验证令牌仅存储 SHA256 摘要。")
    add_body(doc, "用户可加密保存个人 DeepSeek API Key，用于按用户计费与隔离。")

    add_heading3(doc, "3.3.3 大模型与文档处理")
    add_body(doc, "大模型调用通过 OpenAI 兼容 SDK 访问 DeepSeek Chat API；简历与 JD 解析结合 pdfplumber、python-docx 等库；")
    add_body(doc, "图片 JD 场景依赖 Tesseract OCR（含简体中文语言包）。")

    add_screenshot_placeholder(doc, "图3-3 后端 API 文档页面（/docs 或 /health 返回）")

    add_heading2(doc, "3.4 数据存储与中间件")
    add_body(doc, "PostgreSQL 16：存储用户、简历、岗位、匹配结果、对话会话等关系型数据。")
    add_body(doc, "Redis 7：登录/忘记密码限流；岗位列表、岗位详情、用户画像等热点缓存。")

    add_heading2(doc, "3.5 部署与运维")
    add_body(doc, "使用 Docker Compose 编排 db、redis、api、web 四服务；前端生产镜像为多阶段构建（Node 构建 + Nginx 静态托管）。")
    add_body(doc, "已在腾讯云 Ubuntu 服务器完成公网部署验证，支持国内镜像源加速构建。")
    add_screenshot_placeholder(doc, "图3-4 Docker 容器运行状态（docker compose ps 截图）")

    add_heading2(doc, "3.6 技术选型对比（简述）")
    add_body(
        doc,
        "后端亦可选 Django/Flask，本作品选用 FastAPI 以突出异步 IO 与 OpenAPI 自动生成；"
        "数据库亦可选 MySQL，PostgreSQL 对 JSONB 字段（岗位标签、解析结果）支持更好；"
        "大模型亦可选 OpenAI/通义等，DeepSeek 在成本与中文表现上更适合课程演示。",
    )

    # ========== 四、开发过程 ==========
    add_heading1(doc, "四、开发过程")

    add_heading2(doc, "4.1 开发环境与流程")
    add_body(doc, "本地 Windows 开发，使用 Git 版本管理；联调阶段通过 Docker Compose 启动完整依赖栈。")
    add_body(doc, "开发流程遵循：需求分析 → 数据库与接口设计 → 后端 API 实现 → 前端页面联调 → 部署与文档整理。")
    add_screenshot_placeholder(doc, "图4-1 项目目录结构（IDE 或资源管理器截图）")

    add_heading2(doc, "4.2 阶段一：需求分析与设计")
    add_body(doc, "编写需求文档，划分 P0/P1 功能；完成 ER 设计与 REST 接口列表；确定认证方案为邮箱验证 + JWT。")

    add_heading2(doc, "4.3 阶段二：后端核心模块")
    add_heading3(doc, "4.3.1 认证模块")
    add_body(
        doc,
        "实现注册（邮件设密）、登录、忘记密码、令牌校验与完成设密接口；"
        "集成 QQ 邮箱 SMTP 发送验证邮件；开发环境可配置 EXPOSE_DEV_VERIFY_LINK 便于无 SMTP 调试。",
    )
    add_screenshot_placeholder(doc, "图4-2 注册成功/邮件验证弹窗")

    add_heading3(doc, "4.3.2 简历与岗位模块")
    add_body(doc, "简历上传接口保存文件并调用解析服务生成结构化 JSON；岗位库支持 CRUD、筛选与 AI 发现导入。")
    add_screenshot_placeholder(doc, "图4-3 简历上传与侧边栏简历列表")

    add_heading3(doc, "4.3.3 匹配与对话模块")
    add_body(doc, "匹配流水线：SQL 过滤 → 规则预排序 → Top-N 调用 LLM 结构化打分 → 返回推荐列表与报告。")
    add_body(doc, "对话模块维护会话与消息表，支持流式输出（SSE）与 Agent 上下文组装。")
    add_screenshot_placeholder(doc, "图4-4 岗位库匹配结果或单岗分析报告")

    add_heading2(doc, "4.4 阶段三：前端实现")
    add_body(doc, "完成登录/注册/设密/忘记密码页面；主界面 ChatPage 集成侧边栏（简历、API Key、会话列表）；")
    add_body(doc, "岗位库 JobsPage 集成心仪岗位粘贴、AI 搜岗与简历分析入口；ProfilePage、MatchPage 与后端 API 对接。")
    add_screenshot_placeholder(doc, "图4-5 求职意向配置页")

    add_heading2(doc, "4.5 阶段四：部署上线")
    add_body(
        doc,
        "在腾讯云 Ubuntu 安装 Docker，配置镜像加速与 PyPI/npm 国内源；"
        "通过 docker compose 构建并启动；配置安全组放行 8080；"
        "编写 DEPLOY.md 与用户指南文档。",
    )
    add_screenshot_placeholder(doc, "图4-6 云服务器浏览器访问首页（公网 IP:8080）")

    add_heading2(doc, "4.6 开发难点与解决")
    add_body(doc, "（1）国内构建 Docker 镜像超时：配置腾讯云 Docker 镜像与 Dockerfile 内 apt/pip/npm 镜像源。")
    add_body(doc, "（2）邮件验证链接与 CORS：统一配置 APP_PUBLIC_URL 与 CORS_ORIGINS 为实际访问地址。")
    add_body(doc, "（3）LLM 响应慢：列表接口加 Redis 缓存，匹配接口前端展示加载状态，避免重复提交。")

    # ========== 五、测试与优化 ==========
    add_heading1(doc, "五、测试与优化")

    add_heading2(doc, "5.1 测试策略")
    add_body(doc, "采用手工功能测试为主，结合接口文档进行 API 测试；部署环境进行端到端验收。")

    add_heading2(doc, "5.2 功能测试")
    add_heading3(doc, "5.2.1 认证流程")
    add_body(doc, "测试用例：注册新邮箱 → 收信 → 打开链接设密 → 登录成功；重复注册、错误密码、未验证登录等异常分支。")
    add_screenshot_placeholder(doc, "图5-1 登录成功进入主界面")

    add_heading3(doc, "5.2.2 简历与岗位")
    add_body(doc, "测试 PDF/DOCX 上传解析；岗位手动录入与筛选；心仪岗位文本粘贴解析；AI 搜岗（需有效 API Key）。")

    add_heading3(doc, "5.2.3 匹配与对话")
    add_body(doc, "在已配置简历与意向前提下触发岗位库匹配，检查返回分数与说明字段；对话多轮上下文与新建会话。")
    add_screenshot_placeholder(doc, "图5-2 AI 对话回复效果（含 Markdown）")

    add_heading2(doc, "5.3 接口与健康检查")
    add_body(doc, "访问 GET /health 检查数据库与 Redis 连通；通过 /docs 页面测试受保护接口的 Bearer Token 鉴权。")
    add_screenshot_placeholder(doc, "图5-3 /health 或 Swagger 文档截图")

    add_heading2(doc, "5.4 性能与安全优化")
    add_body(doc, "Redis 缓存岗位列表与详情，降低数据库压力；登录失败按邮箱与 IP 计数限流。")
    add_body(doc, "生产环境关闭开发验证链接暴露；JWT_SECRET、数据库密码使用强随机值；.env 不提交版本库。")
    add_body(doc, "前端构建采用 npm ci 与多阶段 Docker 镜像减小体积；Nginx 反向代理统一 /api 入口避免跨域。")

    add_heading2(doc, "5.5 测试结论")
    add_body(
        doc,
        "主要功能路径均可完成：用户可注册登录、维护意向、上传简历、浏览与录入岗位、"
        "获得匹配与优化建议并进行 AI 咨询。已知限制：岗位数据部分为种子数据；"
        "AI 功能依赖外部 API 可用性与网络；OCR 对模糊截图识别率有限。"
        "后续可扩展向量检索、岗位爬虫与管理后台等。",
    )

    add_heading2(doc, "5.6 心得体会（请根据个人情况修改）")
    add_body(
        doc,
        "通过本次大作业，完整经历了从需求到部署的软件工程流程，"
        "加深了对前后端分离、异步 API、容器化部署与大模型应用集成的理解。"
        "（请在此处补充个人真实体会，避免与同学雷同。）",
    )

    # ========== 参考文献 / 附录（可选）==========
    add_heading1(doc, "参考文献")
    add_body(doc, "[1] FastAPI 官方文档. https://fastapi.tiangolo.com/")
    add_body(doc, "[2] React 官方文档. https://react.dev/")
    add_body(doc, "[3] DeepSeek API 文档. https://platform.deepseek.com/api-docs/")
    add_body(doc, "[4] Docker 文档. https://docs.docker.com/")

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    print(f"已生成: {OUT}")


if __name__ == "__main__":
    build()
